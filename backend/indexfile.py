"""Lectura del concentrado de pagos (Excel/CSV) que sirve de índice.

El concentrado no tiene un layout fijo: cambia de mes a mes, suele traer
títulos y filas en blanco antes de la tabla, y los encabezados se escriben
de formas distintas ("PAGO", "IMPORTE", "MONTO"...). Por eso aquí no se
asumen posiciones de columna: se busca la fila de encabezados y se
reconocen sinónimos.
"""

import csv
import io
import re
from decimal import Decimal, InvalidOperation

from text_utils import normalize_key

# Sinónimos aceptados por columna, en orden de preferencia. La comparación
# es por "contiene", sobre el encabezado ya normalizado.
COLUMN_KEYWORDS = {
    "requisicion": [
        "NUMERO DE REQUISICION",
        "NO DE REQUISICION",
        "NUM REQUISICION",
        "REQUISICION",
        "REQUI",
        "REQ",
        "FOLIO",
    ],
    "nombre": [
        "NOMBRE DEL BENEFICIARIO",
        "BENEFICIARIO",
        "NOMBRE",
        "PROVEEDOR",
        "A NOMBRE DE",
    ],
    "banco": [
        "BANCO DE PAGO",
        "BANCO PAGADOR",
        "BANCO ORIGEN",
        "BANCO",
    ],
    "pago": [
        "PAGO",
        "IMPORTE",
        "MONTO",
        "CANTIDAD",
        "TOTAL",
    ],
    # Opcionales: si existen, ayudan a desempatar y a validar.
    "empresa": [
        "EMPRESA",
        "RAZON SOCIAL",
        "CUENTA ORIGEN",
    ],
    "semana": [
        "SEMANA",
        "SEM",
    ],
}

# Columnas mínimas para que el concentrado sirva de índice.
REQUIRED_COLUMNS = ["nombre", "pago"]

MAX_HEADER_SCAN_ROWS = 30


def parse_amount(value) -> Decimal | None:
    """Normaliza importes que pueden venir como número o como texto
    ('$2,600.00', '2600', '2,600.00 MXN')."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"[^\d,.\-]", "", text)
    if not text:
        return None
    # Formato mexicano: la coma es separador de miles.
    text = text.replace(",", "")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _match_column(header: str) -> str | None:
    """Devuelve el nombre lógico de columna para un encabezado, o None."""
    key = normalize_key(header)
    if not key:
        return None
    best: tuple[str, int] | None = None
    for logical, keywords in COLUMN_KEYWORDS.items():
        for keyword in keywords:
            if keyword == key or keyword in key:
                # Gana la coincidencia más larga: "BANCO DE PAGO" debe
                # resolverse como banco, no como pago.
                score = len(keyword)
                if best is None or score > best[1]:
                    best = (logical, score)
    return best[0] if best else None


def _detect_header(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    """Encuentra la fila de encabezados y el índice de cada columna lógica."""
    best: tuple[int, dict[str, int]] | None = None
    best_score = 0
    for row_idx, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for col_idx, cell in enumerate(row):
            if cell is None:
                continue
            logical = _match_column(str(cell))
            if logical and logical not in mapping:
                mapping[logical] = col_idx
        score = len(mapping)
        if all(c in mapping for c in REQUIRED_COLUMNS) and score > best_score:
            best = (row_idx, mapping)
            best_score = score
    return best


def _read_grid(filename: str, data: bytes) -> list[list]:
    """Devuelve el contenido del archivo como una matriz de celdas."""
    lower = (filename or "").lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]

    if lower.endswith(".xls"):
        raise ValueError(
            "El formato .xls (Excel 97-2003) no es compatible. "
            "Ábrelo en Excel y guárdalo como .xlsx o .csv."
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ValueError("Falta la dependencia openpyxl para leer archivos .xlsx") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_index(filename: str, data: bytes) -> dict:
    """Lee el concentrado y devuelve sus filas de pago ya normalizadas.

    Retorna un dict con:
      - `rows`: lista de pagos {requisicion, nombre, banco, pago, ...}
      - `columns`: qué columna del archivo se usó para cada campo
      - `warnings`: avisos para mostrar en la interfaz
    """
    warnings: list[str] = []
    grid = _read_grid(filename, data)

    if not grid:
        raise ValueError("El concentrado está vacío.")

    detected = _detect_header(grid)
    if not detected:
        raise ValueError(
            "No se encontraron las columnas del concentrado. Debe incluir al menos "
            "una columna de nombre/beneficiario y una de pago/importe."
        )

    header_row, mapping = detected
    header_labels = {
        logical: str(grid[header_row][col_idx]).strip()
        for logical, col_idx in mapping.items()
    }

    for optional in ("requisicion", "banco"):
        if optional not in mapping:
            warnings.append(
                f"El concentrado no tiene columna de {optional}; ese dato se captura a mano."
            )

    rows: list[dict] = []
    for row_idx, row in enumerate(grid[header_row + 1 :], start=header_row + 2):
        def cell(logical: str):
            col = mapping.get(logical)
            if col is None or col >= len(row):
                return None
            value = row[col]
            if value is None:
                return None
            text = str(value).strip()
            return text or None

        nombre = cell("nombre")
        pago = parse_amount(row[mapping["pago"]] if mapping["pago"] < len(row) else None)

        # Filas de totales, separadores o vacías: se ignoran en silencio.
        if not nombre or pago is None:
            continue
        if normalize_key(nombre) in {"TOTAL", "TOTALES", "SUMA"}:
            continue

        requisicion = cell("requisicion")
        if requisicion:
            # Excel suele entregar los folios numéricos como '8064686.0'.
            requisicion = re.sub(r"\.0$", "", requisicion)

        rows.append(
            {
                "fila": row_idx,
                "requisicion": requisicion or "",
                "nombre": nombre,
                "banco": cell("banco") or "",
                "pago": str(pago),
                "empresa": cell("empresa") or "",
                "semana": cell("semana") or "",
            }
        )

    if not rows:
        raise ValueError("El concentrado no tiene filas de pago legibles.")

    return {
        "rows": rows,
        "columns": header_labels,
        "header_row": header_row + 1,
        "warnings": warnings,
    }
